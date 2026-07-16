# -*- coding: utf-8 -*-
import json
from pptx import Presentation

json_path = 'D:\\job\\notes_data.json'
ppt_path = 'D:\\job\\智汇人才匹配系统_设计方案.pptx'
out_path = 'D:\\job\\智汇人才匹配系统_设计方案_带备注.pptx'

with open(json_path, 'r', encoding='utf-8') as f:
    notes = json.load(f)

prs = Presentation(ppt_path)
print(f'Opened PPT: {len(prs.slides)} slides')

for i, slide in enumerate(prs.slides):
    ns = slide.notes_slide
    ns.notes_text_frame.text = notes[i]
    print(f'Slide {i+1}: {len(notes[i])} chars')

prs.save(out_path)
print(f'Saved to: {out_path}')
print('Done!')
